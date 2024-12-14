from docx import Document


def main():
    doc = Document("output.docx")

    table = doc.tables[0]
    data = {}

    for row in range(1, 4):
        stat_name = table.cell(row, 0).text[1:]
        stat = table.cell(row, 2).text[1:]

        data[stat_name] = stat
    
    print(f"Данные по ATmega328: {data}")


if __name__ == "__main__":
    main()
    